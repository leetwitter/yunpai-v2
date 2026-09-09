"""M2 BOM/SOP 本地实现行为回归（分片 M2，施工依据 ``_migration/rows-S3.md``）。

覆盖 rows-S3.md 判「需补测试」的 2 条（``onboard_m2_bom_template`` / ``list_m2_runs``）及
其余 M2 工具的端到端行为：

- onboard 模板：history xlsx/csv 列结构与编号前缀 proposal + 模板入库 + 历史行入库
  （product_name 取历史文件名 stem，保证后续按产品名检索可召回）；
- 受控生成：内嵌行规范化/别名映射/去重/默认单位；无依据 → 空行 + note（不伪造）；
- list/get run 过滤与详情（含 artifact_paths kind → path）、跨租户不可见、limit 契约钳制；
- SOP Word 生成（文字流程，不冒充图形流程图）；
- 历史检索：m2_history_lines 打分、canonical BOM 经 **V2 M0 读口**召回、诚实空结果；
- ``run_bom_sop_workflow``（workers.m2_bom）env-gated 落 run 记录。
"""

from __future__ import annotations

import json

import pytest

from yunpai_orchestrator import m2_local
from yunpai_orchestrator.m2_local import (
    LOCAL_HANDLERS,
    M2RunStore,
    generate_bom_controlled,
    generate_m2_sop,
    m2_get_run,
    m2_list_runs,
    onboard_bom_template,
    search_m2_bom_history,
)


def _rule_file(tmp_path, extra: dict | None = None) -> str:
    rules = {
        "column_aliases": {"material_name": ["材料名称", "品名"], "quantity": ["用量"]},
        "duplicate_policy": "reject",
        "default_uom": "PCS",
    }
    if extra:
        rules.update(extra)
    p = tmp_path / "rules.json"
    p.write_text(json.dumps(rules, ensure_ascii=False), encoding="utf-8")
    return str(p)


def _history_csv(tmp_path, name: str = "hist.csv") -> str:
    text = ("物料编码,材料名称,用量,单位\n"
            "W-001,线材,0.5,M\n"
            "W-002,插头,1,PCS\n")
    p = tmp_path / name
    p.write_text(text, encoding="utf-8")
    return str(p)


def _history_xlsx(tmp_path, name: str = "hist.xlsx") -> str:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["料号", "品名", "数量", "单位"])
    ws.append(["X-001", "外壳", 2, "PCS"])
    ws.append(["X-002", "线材", 0.8, "M"])
    p = tmp_path / name
    wb.save(p)
    wb.close()
    return str(p)


def _canonical_bom(db_path, *, tenant_id: str = "T-A", product_code: str = "P-CAN",
                   lines: list[dict] | None = None) -> None:
    """经 M0 后端落一条 canonical bom（publish 后 lifecycle_status=active）。"""
    from yunpai_orchestrator.m0_backend import M0Store

    store = M0Store(str(db_path))
    batch = store.ingest([{
        "entity_type": "bom",
        "product_code": product_code,
        "payload": {"product_code": product_code,
                    "lines": lines or [{"material_code": "C-1", "material_name": "canonical 料",
                                        "quantity": 2}],
                    "attributes": {"source_batch": "b1"}},
    }], tenant_id=tenant_id, task_id="T-M0")
    store.publish(batch["batch_id"], actor="tester")


@pytest.fixture()
def env(tmp_path, monkeypatch):
    db = tmp_path / "m2.sqlite"
    monkeypatch.setenv("YUNPAI_M2_DB", str(db))
    monkeypatch.setenv("YUNPAI_M2_ARTIFACT_DIR", str(tmp_path / "artifacts"))
    # 缺省不接 canonical 事实源（缺省路径可能指向部署库）；各用例需要时自行覆盖。
    monkeypatch.setenv("YUNPAI_M0_DB", str(tmp_path / "m0-absent.sqlite"))
    return M2RunStore(str(db))


@pytest.mark.asyncio
async def test_onboard_proposals_and_templates(env, tmp_path):
    rule = _rule_file(tmp_path)
    csv_p = _history_csv(tmp_path)
    xlsx_p = _history_xlsx(tmp_path)
    out = await onboard_bom_template(
        {"rule_package_path": rule, "history_paths": [csv_p, xlsx_p]},
        {"task_id": "T-ONB", "tenant_id": "T-A"})
    assert isinstance(out["proposals"], list) and len(out["proposals"]) == 2
    csv_proposal = next(p for p in out["proposals"] if p["source_file"].endswith("hist.csv"))
    assert "物料编码" in csv_proposal["columns"]  # 原始列名直报（不猜测规范化）
    assert csv_proposal["row_count"] == 2
    assert csv_proposal["numbering_rule"]["detected_prefix"] == "W-"
    xlsx_proposal = next(p for p in out["proposals"] if p["source_file"].endswith("hist.xlsx"))
    assert xlsx_proposal["numbering_rule"]["detected_prefix"] == "X-"
    # 模板入库 + 历史行入库 + run 记录
    templates = env.list_templates("T-A")
    assert len(templates) >= 2 and templates[0]["columns"]
    lines = env.history_lines("T-A")
    assert len(lines) == 4
    # rows-S3 备注修正：历史行的 product_name 取文件名 stem（不再为空），
    # product_code 仍不猜（保持 ""）。
    assert {row["product_name"] for row in lines} == {"hist", "hist"}
    assert {row["product_code"] for row in lines} == {""}
    runs = env.list_runs("T-A")
    assert any(r["tool"] == "onboard_m2_bom_template" for r in runs)
    # 历史文件不可读 → 显式失败
    with pytest.raises(ValueError, match="不可读"):
        await onboard_bom_template({"rule_package_path": rule,
                                    "history_paths": [str(tmp_path / "nope.xlsx")]},
                                   {"task_id": "T-ONB2", "tenant_id": "T-A"})


@pytest.mark.asyncio
async def test_onboard_run_id_and_tenant_isolation(env, tmp_path):
    """onboard 只写本租户；run_id 口径 m2-{task[-10:]}。"""
    rule = _rule_file(tmp_path)
    await onboard_bom_template(
        {"rule_package_path": rule, "history_paths": [_history_csv(tmp_path)]},
        {"task_id": "T-ONB-9", "tenant_id": "T-A"})
    assert env.history_lines("T-B") == []
    run = env.list_runs("T-A")[0]
    assert run["run_id"] == "m2-T-ONB-9"  # 后 10 位
    assert run["tool"] == "onboard_m2_bom_template" and run["status"] == "done"


@pytest.mark.asyncio
async def test_generate_controlled_embedded_lines(env, tmp_path):
    rule = _rule_file(tmp_path)
    profile = {
        "product_code": "P-1", "product_name": "样品",
        "lines": [
            {"material_code": "M-1", "material_name": "料一", "quantity": "2"},
            {"material_code": "M-1", "material_name": "料一改", "quantity": "1"},  # 重复码
            {"material_code": "M-2", "quantity": 3, "uom": "M"},
        ],
    }
    out = await generate_bom_controlled(
        {"product_profile": profile, "rule_package_path": rule},
        {"task_id": "T-GEN", "tenant_id": "T-A"})
    bom = out["standard_bom"]
    assert len(bom["bom_lines"]) == 2  # duplicate_policy=reject 去重
    first = bom["bom_lines"][0]
    assert first["material_code"] == "M-1" and first["uom"] == "PCS"  # default_uom
    assert bom["rules_applied"]["duplicate_policy"] == "reject"
    assert env.history_lines("T-A") and env.list_runs("T-A")
    run = env.list_runs("T-A")[0]
    assert run["tool"] == "generate_m2_bom_controlled" and run["product_code"] == "P-1"


@pytest.mark.asyncio
async def test_generate_controlled_no_source_note(env, tmp_path):
    rule = _rule_file(tmp_path)
    out = await generate_bom_controlled(
        {"product_profile": {"product_code": "P-X", "product_name": "新品"}, "rule_package_path": rule},
        {"task_id": "T-GX", "tenant_id": "T-A"})
    bom = out["standard_bom"]
    assert bom["bom_lines"] == []          # 无确定性依据不伪造
    assert "模型 C 型" in bom["note"]


@pytest.mark.asyncio
async def test_runs_list_filter_and_get(env):
    env.save_run(tenant_id="T-A", run_id="m2-aaa", tool="x", status="done",
                 order_id="SO-9", product_code="P-1", product_name="样品",
                 summary={"k": 1})
    env.add_artifact("T-A", "m2-aaa", "bom_json", "runtime/x.json")
    rows = await m2_list_runs({"product_code": "P-1"}, {"tenant_id": "T-A"})
    assert rows["data"]["total"] == 1
    assert (await m2_list_runs({"product_code": "P-2"}, {"tenant_id": "T-A"}))["data"]["total"] == 0
    got = await m2_get_run({"run_id": "m2-aaa"}, {"tenant_id": "T-A"})
    assert got["data"]["run"]["summary"] == {"k": 1}
    assert got["data"]["run"]["artifacts"][0]["kind"] == "bom_json"
    with pytest.raises(ValueError, match="not found"):
        await m2_get_run({"run_id": "m2-zzz"}, {"tenant_id": "T-A"})
    with pytest.raises(ValueError, match="not found"):
        await m2_get_run({"run_id": "m2-aaa"}, {"tenant_id": "T-B"})  # 跨租户
    with pytest.raises(ValueError, match="需要 run_id"):
        await m2_get_run({}, {"tenant_id": "T-A"})


@pytest.mark.asyncio
async def test_list_runs_limit_is_clamped_to_contract_bounds(env):
    """契约 limit minimum=1 / maximum=200：handler 必须钳制，不得越界透传。"""
    for i in range(3):
        env.save_run(tenant_id="T-A", run_id=f"m2-{i}", tool="x", status="done")
    assert (await m2_list_runs({"limit": 0}, {"tenant_id": "T-A"}))["data"]["total"] == 1
    assert (await m2_list_runs({"limit": -5}, {"tenant_id": "T-A"}))["data"]["total"] == 1
    assert (await m2_list_runs({"limit": 9999}, {"tenant_id": "T-A"}))["data"]["total"] == 3
    assert (await m2_list_runs({}, {"tenant_id": "T-A"}))["data"]["total"] == 3
    assert (await m2_list_runs({"limit": "abc"}, {"tenant_id": "T-A"}))["data"]["total"] == 3


def test_local_handlers_keys():
    assert set(LOCAL_HANDLERS) == {"onboard_m2_bom_template", "generate_m2_bom_controlled",
                                   "list_m2_runs", "get_m2_run", "generate_m2_sop",
                                   "search_m2_bom_history"}


# ---------- SOP Word ----------

@pytest.mark.asyncio
async def test_generate_sop_docx(env, tmp_path):
    out = await generate_m2_sop(
        {"product_name": "中性线", "part_no": "W-1", "document_no": "SOP-80806-001",
         "bom_items": [{"material_code": "M-1", "material_name": "线材", "quantity": "1", "uom": "M"}],
         "routing_steps": [{"sequence_no": 10, "operation_name": "押出", "equipment_code": "EQ-1"}]},
        {"task_id": "T-SOP", "tenant_id": "T-A"})
    assert out["status"] == "generated"
    from docx import Document

    doc = Document(out["artifacts"]["docx"])
    texts = [p.text for p in doc.paragraphs]
    assert any("产品名称：中性线" in t for t in texts)
    assert out["artifacts"]["document_no"] == "SOP-80806-001"
    # 制品写入 YUNPAI_M2_ARTIFACT_DIR（env 覆盖，缺省 runtime/m2-artifacts）
    assert out["artifacts"]["docx"].startswith(str(tmp_path))
    # run + artifact 记录
    runs = env.list_runs("T-A")
    assert any(r["tool"] == "generate_m2_sop" for r in runs)
    with pytest.raises(ValueError, match="需要"):
        await generate_m2_sop({"product_name": "x"}, {"task_id": "T", "tenant_id": "T-A"})


# ---------- 历史检索 ----------

@pytest.mark.asyncio
async def test_search_history_ranks_matches(env, tmp_path):
    rule = _rule_file(tmp_path)
    profile = {"product_code": "P-A", "product_name": "HDMI 高清线",
               "lines": [{"material_code": "HD-1", "material_name": "线材 HDMI", "quantity": "1"}]}
    await generate_bom_controlled({"product_profile": profile, "rule_package_path": rule},
                                  {"task_id": "T-S1", "tenant_id": "T-A"})
    profile2 = {"product_code": "P-B", "product_name": "USB 线",
                "lines": [{"material_code": "UB-1", "material_name": "线材 USB", "quantity": "1"}]}
    await generate_bom_controlled({"product_profile": profile2, "rule_package_path": rule},
                                  {"task_id": "T-S2", "tenant_id": "T-A"})
    hits = await search_m2_bom_history({"product_name": "HDMI"}, {"tenant_id": "T-A"})
    assert hits["results"] and hits["results"][0]["product_code"] == "P-A"
    assert hits["results"][0]["score"] > 0
    assert hits["results"][0]["lines_count"] >= 1
    assert hits["warnings"] == []  # canonical 库不存在 = 没有事实源，不算降级
    kw = await search_m2_bom_history({"product_name": "x", "keywords": ["UB-1"]},
                                     {"tenant_id": "T-A"})
    assert kw["results"] and any(r["product_code"] == "P-B" for r in kw["results"])
    none = await search_m2_bom_history({"product_name": "不存在"}, {"tenant_id": "T-A"})
    assert none["results"] == []  # 诚实空
    assert env.history_lines("T-A")


@pytest.mark.asyncio
async def test_search_reads_canonical_bom_through_m0_read_port(env, tmp_path, monkeypatch):
    """rows-S3 第 1 行改造项：canonical 候选必须经 V2 M0 读口拿到，不再直连 sqlite 表。"""
    m0_db = tmp_path / "m0.sqlite"
    monkeypatch.setenv("YUNPAI_M0_DB", str(m0_db))
    _canonical_bom(m0_db, tenant_id="T-A", product_code="P-CAN",
                   lines=[{"material_code": "C-1", "material_name": "canonical 料", "quantity": 2}])
    out = await search_m2_bom_history({"product_name": "P-CAN"}, {"tenant_id": "T-A"})
    assert out["warnings"] == []
    hit = next((r for r in out["results"] if r["source"] == "canonical_bom"), None)
    assert hit is not None, "canonical BOM 候选未被 M0 读口读出"
    assert hit["product_code"] == "P-CAN" and hit["lines_count"] == 1
    assert hit["lines"][0]["material_code"] == "C-1"
    # 跨租户不可见
    other = await search_m2_bom_history({"product_name": "P-CAN"}, {"tenant_id": "T-B"})
    assert all(r["source"] != "canonical_bom" for r in other["results"])


@pytest.mark.asyncio
async def test_search_reports_read_port_failure_in_warnings(env, tmp_path, monkeypatch):
    """读口异常不得静默退化成空候选：必须写进 warnings。"""
    broken = tmp_path / "broken-m0.sqlite"
    broken.write_text("not a sqlite database", encoding="utf-8")
    monkeypatch.setenv("YUNPAI_M0_DB", str(broken))
    out = await search_m2_bom_history({"product_name": "任意"}, {"tenant_id": "T-A"})
    assert out["results"] == []
    assert out["warnings"] and "canonical BOM 读口不可用" in out["warnings"][0]


@pytest.mark.asyncio
async def test_search_consumes_declared_history_paths(env, tmp_path):
    hist = _history_csv(tmp_path, "hist-file.csv")
    out = await search_m2_bom_history({"product_name": "hist-file", "history_paths": [hist]},
                                      {"tenant_id": "T-A"})
    assert out["results"], "history_paths 候选未被检索"
    hit = out["results"][0]
    assert hit["source"] == "history_file" and hit["lines_count"] == 2
    with pytest.raises(ValueError, match="不可读"):
        await search_m2_bom_history(
            {"product_name": "x", "history_paths": [str(tmp_path / "nope.csv")]},
            {"tenant_id": "T-A"})


def test_canonical_body_normalizes_attributes_top_level_wins():
    """V2 无 fact_gateway：``payload.attributes`` 并入顶层，顶层优先。"""
    body = m2_local._canonical_body(
        {"payload": {"role": "sop", "attributes": {"route_steps": [{"sequence_no": 10}], "title": "T"}}})
    assert body["route_steps"] == [{"sequence_no": 10}] and body["title"] == "T"
    assert m2_local._canonical_body(
        {"payload": {"lines": ["top"], "attributes": {"lines": ["attr"]}}})["lines"] == ["top"]


# ---------- run_bom_sop_workflow 联动 ----------

@pytest.mark.asyncio
async def test_run_bom_sop_workflow_records_run_when_env_set(env):
    """workers.m2_bom 在配置 YUNPAI_M2_DB 时落 run 记录（env-gated 联动）。"""
    from yunpai_orchestrator.workers import m2_bom

    payload = {"product_profile": {"product_code": "P-WF", "product_name": "联动产品"},
               "bom_lines": [{"material_code": "M-1", "material_name": "料", "quantity": 1, "uom": "PCS"}]}
    out = await m2_bom(payload, {"task_id": "T-WF", "tenant_id": "T-A"})
    assert out["status"] == "draft_created"
    runs = env.list_runs("T-A")
    run = next(r for r in runs if r["tool"] == "run_bom_sop_workflow")
    assert run["run_id"] == "m2-T-WF"
    assert run["product_code"] == "P-WF" and run["status"] == "draft_created"
