from io import BytesIO

from openpyxl import Workbook

from yunpai_orchestrator.order_workbook import parse_order_workbook


def test_parse_stocking_order_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet["P6"] = "PO-TEST-001"
    sheet["X6"] = "2026年8月12日"
    sheet["X7"] = "2026-09-01"
    sheet["G6"] = "供应商"
    sheet["E9"] = "序号"
    sheet["I9"] = "型号"
    sheet["J9"] = "名称"
    sheet["R9"] = "采购数量"
    sheet["V9"] = "单价"
    sheet["W9"] = "金额"
    sheet["E10"] = 1
    sheet["I10"] = "W-H909"
    sheet["J10"] = "高清线"
    sheet["R10"] = 4000
    sheet["V10"] = 3.4
    sheet["W10"] = 13600
    output = BytesIO()
    workbook.save(output)
    document = parse_order_workbook("order.xlsx", output.getvalue())
    assert document["order_id"] == "PO-TEST-001"
    assert document["product_code"] == "W-H909"
    assert document["quantity"] == 4000
    assert document["lines"][0]["amount"] == 13600


def test_parse_order_workbook_tolerates_text_packaging_and_placeholders():
    workbook = Workbook()
    sheet = workbook.active
    sheet["P6"] = "PO-TEXT-001"
    sheet["X7"] = "2026-09-01"
    sheet["E10"], sheet["I10"], sheet["R10"] = 1, "FC-35", 24
    sheet["V10"], sheet["W10"] = 50, 1200
    sheet["Z10"] = "/"
    sheet["AB10"] = "中性彩盒+彩盒贴LOGO贴纸"
    sheet["AE10"], sheet["AF10"] = 12, 2
    output = BytesIO()
    workbook.save(output)

    document = parse_order_workbook("real-order.xlsx", output.getvalue())

    assert document["quantity"] == 24
    assert document["lines"][0]["packaging"] == "中性彩盒+彩盒贴LOGO贴纸"
    assert document["lines"][0]["factory_inventory"] is None
    assert document["lines"][0]["pack_quantity"] == 12
    assert document["lines"][0]["carton_count"] == 2
    assert document["validation_issues"] == []
