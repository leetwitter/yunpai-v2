"""图片/PDF 多模态采样（确定性读内容）。"""

from __future__ import annotations

import base64
import io

from yunpai_orchestrator.media_sample import downscale_image, render_pdf_pages
from yunpai_orchestrator.recognized_store import sample_file


def _png() -> bytes:
    from PIL import Image

    img = Image.new("RGB", (400, 200), "white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _pdf() -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "yunpai test")
    raw = doc.tobytes()
    doc.close()
    return raw


def test_downscale_image_returns_jpeg_base64():
    b64 = downscale_image(_png())
    raw = base64.b64decode(b64)
    assert raw[:3] == b"\xff\xd8\xff"  # JPEG magic


def test_sample_file_png_returns_images():
    sample = sample_file(_png(), "x.png")
    assert sample["sniff"]["detected_format"] == "png"
    assert len(sample["images"]) == 1
    assert sample["content_sampled"] is True


def test_render_pdf_pages_returns_images():
    images = render_pdf_pages(_pdf(), max_pages=1)
    assert len(images) == 1
    assert base64.b64decode(images[0])[:3] == b"\xff\xd8\xff"


def test_sample_file_pdf_returns_images():
    sample = sample_file(_pdf(), "x.pdf")
    assert sample["sniff"]["detected_format"] == "pdf"
    assert len(sample["images"]) >= 1
    assert sample["content_sampled"] is True


def test_sample_file_unsupported_image_degrades_gracefully():
    # 伪 png（内容不是真图片）→ 嗅探为 png 但渲染失败，降级为无 images，不抛异常。
    sample = sample_file(b"\x89PNG\r\n\x1a\n" + b"garbage", "fake.png")
    assert sample["sniff"]["detected_format"] == "png"
    assert sample["images"] == []


def test_sample_file_multi_sheet_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "材料明细"
    ws1.append(["物料编码", "材料名称", "规格"])
    ws1.append(["YA.001", "铜箔", "2m"])
    ws2 = wb.create_sheet("产品BOM-1")
    ws2.append(["物料编码", "材料名称", "用量"])
    ws2.append(["YA.001", "铜箔", 2])
    ws3 = wb.create_sheet("产品BOM-2")
    ws3.append(["物料编码", "材料名称", "用量"])
    ws3.append(["YA.002", "线材", 5])
    buf = io.BytesIO()
    wb.save(buf)

    sample = sample_file(buf.getvalue(), "BOM.xlsx", max_sheets=3)
    assert sample["sheet_names"] == ["材料明细", "产品BOM-1", "产品BOM-2"]
    assert len(sample["sheets"]) == 3
    assert sample["sheets"][0]["name"] == "材料明细"
    assert sample["sheets"][0]["headers"][:3] == ["物料编码", "材料名称", "规格"]
    # 主 headers 取第一个有内容的 sheet
    assert sample["headers"][:3] == ["物料编码", "材料名称", "规格"]
