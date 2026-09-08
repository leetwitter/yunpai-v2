"""order_semantics 语义补充层 + 真实订单结构同构复现测试。

真实样本含客户数据，不能进入仓库（原始文件与重放证据保存在仓库外的
run 目录）；本文件用**同构合成工作簿**（相同的单 Sheet / 块事实 / 表头行 /
明细行 / 合计行 / 合并备注区布局，数值与名称全部替换为虚构值）来固化解析
语义。真实文件的本地重放证据单独记录在 run 目录（仓库外）。
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from yunpai_orchestrator.order_parser_v2 import parse_order_sheets
from yunpai_orchestrator.order_semantics import (
    build_semantic_supplement,
    parse_order_document,
    result_has_order_gap,
    workbook_looks_like_order,
)

XLSX_SUFFIX = ".xlsx"


def real_order_twin_bytes() -> bytes:
    """复刻真实订单同构布局：表头行 7、块事实 1-6、明细 8-9、合计 10、
    合并备注区 13（“付款流水…”文本行绝不能成为订单行）。

    列：产品图片|系统型号|名称/材质要求|规格|采购数量|销售价|装箱数量|件数|金额|
        备注(标签尺寸)|标签数量|备注(标签LOGO)…
    明细行系统型号列值为“无”（无编码占位）。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1"
    sheet["A1"] = "业务订单"
    sheet["A2"] = "采购订单号："
    sheet["B2"] = "WX20260905001"
    sheet["A3"] = "客户："
    sheet["B3"] = "测试客户A"
    sheet["A4"] = "联系方式："
    sheet["B4"] = "13800000000"
    sheet["A5"] = "订购日期："
    sheet["B5"] = datetime(2026, 9, 1, 0, 0, 0)
    sheet["A6"] = "交货日期："
    sheet["A7"] = "产品图片"
    sheet["B7"] = "系统型号"
    sheet["C7"] = "名称/材质要求"
    sheet["D7"] = "规格"
    sheet["E7"] = "采购数量"
    sheet["F7"] = "销售价"
    sheet["G7"] = "装箱数量"
    sheet["H7"] = "件数"
    sheet["I7"] = "金额"
    sheet["J7"] = "备注(标签尺寸)"
    sheet["B8"] = "无"
    sheet["C8"] = "测试线材 型号V1 OD5.0"
    sheet["D8"] = "15M"
    sheet["E8"] = 500
    sheet["F8"] = 60
    sheet["G8"] = 12
    sheet["I8"] = 30000
    sheet["B9"] = "无"
    sheet["D9"] = "20M"
    sheet["E9"] = 500
    sheet["F9"] = 62
    sheet["G9"] = 12
    sheet["I9"] = 31000
    sheet["H10"] = "合计"
    sheet["I10"] = 61000
    sheet["A13"] = "付款流水请截图以下"
    sheet.merge_cells("A13:C13")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def inventory_like_bytes() -> bytes:
    """库存登记表布局（仓库/物料编码/现存数量）—— 不得被改判为订单。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "库存"
    sheet.append(["仓库", "物料编码", "物料名称", "现存数量", "库位"])
    sheet.append(["成品仓", "M-1001", "测试线材", 120, "A-01"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def coded_order_twin_bytes() -> bytes:
    """同构订单但带有可提升为权威产品编码的系统型号。"""
    workbook = load_workbook(BytesIO(real_order_twin_bytes()))
    sheet = workbook.active
    sheet["B8"] = "W-H909"
    sheet["B9"] = "W-H911"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parser_produces_valid_header_lines_and_review_issues_for_real_layout():
    raw = real_order_twin_bytes()
    document = parse_order_sheets("订单-同构-回放.xlsx", raw)
    assert document["order_id"] == "WX20260905001"
    assert document["order_date"] == "2026-09-01"
    assert document["customer_name"] == "测试客户A"
    assert document["line_count"] == 2
    assert document["total_quantity"] == 1000.0
    assert document["total_amount"] == 61000.0
    lines = document["lines"]
    assert [line["row"] for line in lines] == [8, 9]
    assert all(line["product_code"] is None for line in lines)  # “无”不得冒充编码
    assert lines[0]["product_name"].startswith("测试线材")
    assert lines[0]["specification"] == "15M"
    assert lines[0]["quantity"] == 500.0
    assert lines[0]["unit_price"] == 60.0
    assert lines[0]["amount"] == 30000.0
    assert lines[0]["pack_quantity"] == 12.0
    assert any(issue["code"] == "MISSING_PRODUCT_CODE" for issue in document["validation_issues"])
    assert any(item and item["code"] == "MISSING_ORDER_NUMBER" for item in document["review_issues"]) is False
    assert any(item and item["code"] == "MISSING_LINE_PRODUCT_CODE" for item in document["review_issues"])
    assert document["due_date"] == ""


def test_semantic_supplement_for_external_order_gap_keeps_both_results():
    raw = real_order_twin_bytes()
    external = {
        "task_id": "run-5b4eeb-real-shape",
        "status": "needs_review",
        "processing_stage": "review",
        "needs_review": True,
        "overall_confidence": 0.4,
        "doc_type": "order",
        "document_subtype": "customer_purchase_order",
        "document": {
            "schema_version": "m1.document.v2",
            "document_type": "order",
            "document_kind_label": "customer_purchase_order",
            "header": {
                "order_number": None, "doc_date": None, "title": "业务订单",
                "order_id": None, "customer": None, "supplier": None,
            },
            "lines": [],
        },
    }
    assert result_has_order_gap(external) is True
    supplement = build_semantic_supplement("order.xlsx", raw, external=external)
    assert supplement is not None
    assert supplement["schema_version"] == "m1.semantic-supplement.v1"
    assert supplement["requires_review"] is True
    assert supplement["document"]["header"]["order_number"] == "WX20260905001"
    assert supplement["document"]["header"]["customer_name"] == "测试客户A"
    assert len(supplement["document"]["lines"]) == 2
    assert supplement["source"]["sha256"]
    assert any(missing["field"] == "line.product_code" for missing in supplement["missing_fields"])
    assert any(conflict["field"] == "lines" for conflict in supplement["conflicts"])
    # 外部结果必须原样保留（本测试只验证补充，不覆盖外部）。
    assert external["document"]["lines"] == []


def test_semantic_supplement_repairs_missing_header_product_code():
    raw = coded_order_twin_bytes()
    external = {
        "status": "needs_review", "doc_type": "order",
        "document": {
            "document_type": "order", "document_subtype": "stocking_order",
            "header": {"order_number": "PO-1"},
            "lines": [{"model": "W-H909", "product_code": "W-H909"}],
        },
    }
    assert result_has_order_gap(external) is True
    supplement = build_semantic_supplement("order.xlsx", raw, external=external)
    assert supplement is not None
    assert supplement["document"]["header"]["product_code"] == "W-H909"
    assert not any(item["field"] == "header.product_code" for item in supplement["missing_fields"])


def test_semantic_supplement_can_promote_external_order_line_model():
    external = {
        "status": "needs_review", "doc_type": "order",
        "document": {
            "header": {"order_number": "PO-1"},
            "lines": [{"model": "W-H909", "product_code": None}],
        },
    }
    supplement = build_semantic_supplement("order.xlsx", real_order_twin_bytes(), external=external)
    assert supplement is not None
    assert supplement["document"]["header"]["product_code"] == "W-H909"


def test_no_supplement_for_non_order_or_complete_external():
    assert result_has_order_gap({"task_id": "t", "status": "done"}) is False
    complete = {
        "task_id": "t", "status": "done", "doc_type": "order",
        "document": {"document_type": "order", "header": {"order_number": "SO-1"}, "lines": [{"line_id": "L1"}]},
    }
    assert result_has_order_gap(complete) is False


def test_inventory_table_is_not_reclassified_as_order():
    raw = inventory_like_bytes()
    assert workbook_looks_like_order(raw, "库存登记-202609.xlsx") is False
    assert workbook_looks_like_order(real_order_twin_bytes(), "备货订单-同构.xlsx") is True


def test_parse_order_document_reports_parser_version_and_sha():
    raw = real_order_twin_bytes()
    parsed = parse_order_document("订单-同构.xlsx", raw)
    assert parsed["parser_version"]
    assert parsed["source"]["sha256"]
    assert parsed["document"]["order_id"] == "WX20260905001"
    assert len(parsed["document"]["lines"]) == 2
