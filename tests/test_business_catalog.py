from __future__ import annotations

import json

import pytest
from openpyxl import Workbook

from yunpai_orchestrator.business_catalog import (
    catalog_summary,
    ingest_tree,
    list_candidates,
    transition_candidate_review,
)


def _order_book(path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["P6"] = "PO-TEST-001"
    sheet["X6"] = "2026-09-01"
    sheet["X7"] = "2026-09-10"
    sheet["G6"] = "供应商"
    sheet["P7"] = "月结"
    sheet["D8"] = "东莞"
    sheet["E10"] = 1
    sheet["I10"] = "W-H909"
    sheet["J10"] = "HDMI"
    sheet["N10"] = "4K"
    sheet["Q10"] = "1M"
    sheet["R10"] = 4
    sheet["S10"] = "PCS"
    sheet["V10"] = 10
    sheet["W10"] = 40
    workbook.save(path)


def test_ingest_tree_extracts_order_and_is_repeatable(tmp_path):
    root = tmp_path / "业务数据"
    root.mkdir()
    _order_book(root / "桐曦备货订单.xlsx")
    (root / "notes.json").write_text(json.dumps({"records": [{"kind": "bom"}]}), encoding="utf-8")
    db = tmp_path / "catalog.sqlite"

    first = ingest_tree(root, db, parse_xlsx=True)
    second = ingest_tree(root, db, parse_xlsx=True)

    assert first["file_count"] == 2
    assert first["error_count"] == 0
    assert second["error_count"] == 0
    summary = catalog_summary(db)
    assert summary["source_files"] == 2
    assert summary["documents"] == 2
    assert summary["field_observations"] > 0


def test_ingest_tree_extracts_all_bom_sheets_and_rows(tmp_path):
    root = tmp_path / "BOM专项"
    root.mkdir()
    workbook = Workbook()
    first = workbook.active
    first.title = "产品A"
    first.append(["项目", "物料编码", "材料名称", "用量", "单位", "单价", "成本"])
    first.append([1, "YA.A.01.001", "插头料", 2, "PCS", 1.2, 2.4])
    second = workbook.create_sheet("产品B")
    second.append(["物料编码", "材料名称", "规格", "数量", "单位"])
    second.append(["XC005", "光纤", "30#", 1, "M"])
    workbook.save(root / "成品BOM多Sheet.xlsx")

    db = tmp_path / "catalog.sqlite"
    result = ingest_tree(root, db, parse_xlsx=True)

    assert result["error_count"] == 0
    summary = catalog_summary(db)
    assert summary["field_observations"] >= 9
    import sqlite3
    with sqlite3.connect(db) as connection:
        payload = connection.execute("SELECT payload_json FROM document_candidates").fetchone()[0]
        metadata = connection.execute("SELECT metadata_json FROM source_files").fetchone()[0]
    assert '"sheet_count": 2' in payload
    assert '"bom_line_count": 2' in payload
    assert '"产品A"' in metadata and '"产品B"' in metadata


def test_ingest_tree_keeps_binary_business_file_types(tmp_path):
    root = tmp_path / "BOM规则"
    root.mkdir()
    (root / "verify_release.py").write_text("print('ok')", encoding="utf-8")
    result = ingest_tree(root, tmp_path / "catalog.sqlite")
    assert result["file_count"] == 1
    assert result["error_count"] == 0
    assert result["counts_by_kind"] == {"bom": 1}


def test_content_classifier_upgrades_unnamed_equipment_sheet(tmp_path):
    from openpyxl import Workbook

    root = tmp_path / "资料"
    root.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["设备编码", "设备名称", "规格", "产线", "能力"])
    sheet.append(["EQ-001", "注塑机", "HTF250", "线1", "250T"])
    workbook.save(root / "无名表.xlsx")

    db = tmp_path / "catalog.sqlite"
    result = ingest_tree(root, db, parse_xlsx=True)
    assert result["error_count"] == 0
    assert result["counts_by_kind"].get("equipment", 0) == 1
    import sqlite3

    with sqlite3.connect(db) as connection:
        row = connection.execute("SELECT review_status, sensitivity_classification, classification_confidence, document_type FROM document_candidates").fetchone()
    assert row[0] == "candidate" or row[0] == "needs_review"
    assert row[1] == "internal"
    assert row[3] == "equipment"


def test_missing_fields_force_needs_review(tmp_path):
    root = tmp_path / "残缺BOM"
    root.mkdir()
    # 只有材料名称无编码/用量 -> bom 最低字段不足。
    (root / "bom残缺.xlsx").write_bytes(_bom_with_only_name())
    db = tmp_path / "catalog.sqlite"
    result = ingest_tree(root, db, parse_xlsx=True)
    assert result["error_count"] == 0
    import sqlite3

    with sqlite3.connect(db) as connection:
        payload = connection.execute("SELECT payload_json, missing_fields_json, review_status FROM document_candidates").fetchone()
    assert payload[2] == "needs_review"
    assert '"material_code"' in payload[1]


def test_hr_sensitivity_marks_worker_and_wage_paths(tmp_path):
    root = tmp_path / "人事资料"
    root.mkdir()
    (root / "员工工资表.xlsx").write_bytes(_bom_with_only_name())
    db = tmp_path / "catalog.sqlite"
    ingest_tree(root, db, parse_xlsx=True)
    import sqlite3

    with sqlite3.connect(db) as connection:
        sensitivity = connection.execute("SELECT sensitivity_classification FROM document_candidates").fetchone()[0]
    assert sensitivity == "hr"


def _bom_with_only_name() -> bytes:
    from io import BytesIO

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["材料名称"])
    sheet.append(["普通螺丝"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_candidate_review_state_machine_records_actor(tmp_path):
    root = tmp_path / "BOM审批"
    root.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["物料编码", "材料名称", "用量"])
    sheet.append(["YA.A.01.001", "插头料", 2])
    workbook.save(root / "bom.xlsx")
    db = tmp_path / "catalog.sqlite"
    ingest_tree(root, db, parse_xlsx=True)
    candidates = list_candidates(db)
    assert len(candidates) == 1
    document_id = candidates[0]["document_id"]
    assert candidates[0]["review_status"] == "candidate"

    approved = transition_candidate_review(db, document_id, decision="approve", actor="zhb", entity_key="product:YA.A.01.001", entity_version="v1")
    assert approved["from_status"] == "candidate"
    assert approved["to_status"] == "approved"
    assert approved["reviewer"] == "zhb"
    listed = list_candidates(db, review_status="approved")[0]
    assert listed["entity_version"] == "v1"
    assert listed["reviewer"] == "zhb"
    assert listed["reviewed_at"]

    with pytest.raises(ValueError):
        transition_candidate_review(db, document_id, decision="approve", actor="zhb")


def test_candidate_review_rejects_unknown_and_bad_decision(tmp_path):
    db = tmp_path / "catalog.sqlite"
    with pytest.raises(ValueError):
        transition_candidate_review(db, "doc-nope", decision="approve", actor="zhb")
    root = tmp_path / "bad"
    root.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["物料编码", "材料名称", "用量"])
    sheet.append(["YA.A.01.001", "插头料", 2])
    workbook.save(root / "bom.xlsx")
    ingest_tree(root, db, parse_xlsx=True)
    document_id = list_candidates(db)[0]["document_id"]
    with pytest.raises(ValueError):
        transition_candidate_review(db, document_id, decision="nonsense")


def test_large_xlsx_is_marked_deferred_not_completed(tmp_path):
    """大表延迟深解析：状态必须显式 deferred_to_m1，不能声称已识别完成。"""
    import sqlite3

    root = tmp_path / "大表"
    root.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["物料编码", "材料名称", "用量"])
    for index in range(5):
        sheet.append([f"M-{index}", "材料", index + 1])
    workbook.save(root / "big-bom.xlsx")
    db = tmp_path / "deferred.sqlite"
    # parse_xlsx=False 且超过 deep 阈值走轻量登记 -> deferred。
    result = ingest_tree(root, db, parse_xlsx=False, deep_limit_bytes=1)
    with sqlite3.connect(db) as connection:
        row = connection.execute("SELECT status FROM source_files").fetchone()
        doc = connection.execute("SELECT review_status FROM document_candidates").fetchone()
    assert row[0] == "deferred_to_m1"
    assert doc[0] == "deferred_to_m1"
