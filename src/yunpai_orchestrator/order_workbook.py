from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any


_NULL_NUMERIC_VALUES = {"", "/", "-", "--", "—", "n/a", "na", "无", "none", "null"}


def _value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str):
        return value.strip()
    return value


def _number(value: Any, *, issues: list[dict[str, Any]], field: str, row: int, default: float | None = None) -> float | None:
    """Coerce spreadsheet numbers without making a whole order unreadable."""
    value = _value(value)
    if value is None:
        return default
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text.lower() in _NULL_NUMERIC_VALUES:
        return default
    try:
        return float(text.replace(",", ""))
    except (TypeError, ValueError):
        issues.append({"code": "INVALID_NUMBER", "field": field, "row": row, "raw_value": text})
        return default


def _text_candidate(sheet: Any, row: int, columns: range) -> str:
    """Find packaging text across workbook revisions that shift image columns."""
    for column in columns:
        value = _value(sheet.cell(row, column).value)
        if isinstance(value, str) and value and not value.startswith("="):
            try:
                float(value)
            except ValueError:
                return value
    return ""


def _numeric_candidates(sheet: Any, row: int, columns: range, issues: list[dict[str, Any]], field_prefix: str, *, skip_values: set[str] | None = None) -> list[float]:
    values: list[float] = []
    for column in columns:
        value = _value(sheet.cell(row, column).value)
        if value is None or (isinstance(value, str) and (not value.strip() or value.strip().startswith("="))):
            continue
        if isinstance(value, str) and value.strip() in (skip_values or set()):
            continue
        number = _number(value, issues=issues, field=f"{field_prefix}_{column}", row=row)
        if number is not None:
            values.append(number)
    return values


def parse_order_workbook(filename: str, raw: bytes) -> dict[str, Any]:
    """Extract the stable order facts from the provided stocking-order workbook."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.datetime import from_excel
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("XLSX upload requires openpyxl") from exc
    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"invalid XLSX file: {exc}") from exc
    sheet = workbook.active
    cells = {cell.coordinate: _value(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None}

    def cell(address: str, default: Any = None) -> Any:
        value = cells.get(address, default)
        if isinstance(value, (int, float)) and address in {"X7"}:
            try:
                return from_excel(value, workbook.epoch).date().isoformat()
            except Exception:
                return value
        return value

    lines: list[dict[str, Any]] = []
    validation_issues: list[dict[str, Any]] = []
    for row in range(10, sheet.max_row + 1):
        sequence = sheet.cell(row, 5).value
        model = _value(sheet.cell(row, 9).value)
        quantity = sheet.cell(row, 18).value
        if sequence in (None, "") or not model or quantity in (None, ""):
            continue
        packaging = _text_candidate(sheet, row, range(28, 31))
        packaging_value = _value(sheet.cell(row, 28).value)
        numeric_tail = _numeric_candidates(sheet, row, range(30, 34), validation_issues, "tail", skip_values={packaging} if packaging else set())
        lines.append({
            "line_id": f"{str(cells.get('P6') or 'order')}-{sequence}",
            "line_no": int(sequence) if str(sequence).isdigit() else sequence,
            "model": str(model),
            "product_code": str(model),
            "product_name": str(_value(sheet.cell(row, 10).value) or ""),
            "name_raw": str(_value(sheet.cell(row, 10).value) or ""),
            "specification": str(_value(sheet.cell(row, 14).value) or ""),
            "specification_raw": str(_value(sheet.cell(row, 14).value) or ""),
            "size": str(_value(sheet.cell(row, 17).value) or ""),
            "quantity": _number(quantity, issues=validation_issues, field="quantity", row=row),
            "unit": str(_value(sheet.cell(row, 19).value) or "PCS"),
            "unit_price": _number(sheet.cell(row, 22).value, issues=validation_issues, field="unit_price", row=row, default=0),
            "amount": _number(sheet.cell(row, 23).value, issues=validation_issues, field="amount", row=row, default=0),
            "unit_price_tax_included": _number(sheet.cell(row, 22).value, issues=validation_issues, field="unit_price_tax_included", row=row, default=0),
            "amount_tax_included": _number(sheet.cell(row, 23).value, issues=validation_issues, field="amount_tax_included", row=row, default=0),
            "dongguan_inventory": _number(sheet.cell(row, 25).value, issues=validation_issues, field="dongguan_inventory", row=row),
            "factory_inventory": _number(sheet.cell(row, 26).value, issues=validation_issues, field="factory_inventory", row=row),
            "in_transit": _number(sheet.cell(row, 27).value, issues=validation_issues, field="in_transit", row=row),
            "three_month_sales": None if packaging and packaging_value == packaging else _number(sheet.cell(row, 28).value, issues=validation_issues, field="three_month_sales", row=row),
            "packaging": packaging,
            "pack_quantity": numeric_tail[0] if numeric_tail else None,
            "carton_count": numeric_tail[1] if len(numeric_tail) > 1 else None,
            "remark": str(_value(sheet.cell(row, 34).value) or ""),
        })
    total_quantity = sum(line["quantity"] or 0 for line in lines)
    total_amount = sum(line["amount"] or 0 for line in lines)
    return {
        "order_id": str(cell("P6") or ""),
        "order_date": str(cell("X6") or ""),
        "due_date": str(cell("X7") or ""),
        "supplier_name": str(cell("G6") or ""),
        "payment_terms": str(cell("P7") or ""),
        "delivery_address": str(cell("D8") or ""),
        "product_code": lines[0]["product_code"] if lines else None,
        "quantity": total_quantity,
        "total_amount": total_amount,
        "lines": lines,
        "document_type": "order",
        "document_subtype": "stocking_order",
        "confidence": 0.98 if lines and cell("P6") and cell("X7") and not validation_issues else 0.78 if lines else 0.65,
        "validation_issues": validation_issues,
        "source_filename": filename,
    }
